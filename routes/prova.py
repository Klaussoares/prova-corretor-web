import os
import io
import json
import datetime as dt
from datetime import datetime
import tempfile
import shutil
import logging
from werkzeug.utils import secure_filename
from pdf2image import convert_from_path, pdfinfo_from_path
from flask import Blueprint, request, jsonify, send_file
from openpyxl import Workbook, load_workbook
from flask_cors import cross_origin
from src.config import (
    EMAILS_AUTORIZADOS,
    GEMINI_API_KEY,
    COORDS_NOVO,
    COORDS_ANTIGO,
    CLEANUP_DELAY_SECONDS,
    LOG_DIR,
    LOG_FILE,
    is_email_authorized,
    get_log_entry
)

import google.generativeai as genai

if GEMINI_API_KEY:
    genai.configure(api_key=GEMINI_API_KEY)
    model = genai.GenerativeModel("gemini-1.5-flash")
else:
    model = None

DPI = int(os.environ.get("PDF_DPI", "150"))


prova_bp = Blueprint('prova', __name__)

# Imports condicionais para evitar erros no deploy
# model = None
convert_from_path = None
Image = None
cv2 = None
np = None
pd = None

try:
    if GEMINI_API_KEY:
        import google.generativeai as genai
        genai.configure(api_key=GEMINI_API_KEY)
        model = genai.GenerativeModel(
            model_name="gemini-2.0-flash",
            system_instruction="Você é excelente ajudante para corrigir rapidamente provas de alunos."
        )
except ImportError:
    logging.warning("Google Generative AI não disponível no ambiente de deploy")

try:
    from pdf2image import convert_from_path
    from PIL import Image
    import cv2
    import numpy as np
    import pandas as pd
except ImportError:
    logging.warning("Dependências de processamento de imagem não disponíveis no ambiente de deploy")

# Configuração de logging
os.makedirs(LOG_DIR, exist_ok=True)
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler(LOG_FILE, encoding='utf-8'),
        logging.StreamHandler()
    ]
)

def log_action(email, action, details=None, error=None):
    """Registra uma ação no log."""
    log_entry = get_log_entry(email, action, details, error)
    
    # Log no arquivo
    with open(LOG_FILE, 'a', encoding='utf-8') as f:
        f.write(log_entry + '\n')
    
    # Log no console
    if error:
        logging.error(log_entry)
    else:
        logging.info(log_entry)

def crop_pil(img, box):
    """Corta uma imagem PIL usando as coordenadas fornecidas."""
    return img.crop(box)

def preprocessar_para_ia(imagem_pil):
    """Preprocessa a imagem para melhor reconhecimento pela IA."""
    img_cv = cv2.cvtColor(np.array(imagem_pil), cv2.COLOR_RGB2BGR)
    img_gray = cv2.cvtColor(img_cv, cv2.COLOR_BGR2GRAY)
    return Image.fromarray(cv2.cvtColor(img_gray, cv2.COLOR_GRAY2RGB))

def carregar_gabaritos_excel(excel_file):
    """Carrega os gabaritos do arquivo Excel fornecido."""
    gabaritos = {}
    try:
        wb_gab = load_workbook(excel_file)
        for sheet_name in wb_gab.sheetnames:
            if "modelo" in sheet_name.lower():
                numero = ''.join(filter(str.isdigit, sheet_name))
                ws_gab = wb_gab[sheet_name]
                respostas = [str(r[1]).strip().upper() for r in ws_gab.iter_rows(min_row=2, max_col=2, values_only=True)
                             if r[1] and str(r[1]).strip().upper() in ['A', 'B', 'C', 'D', 'E']]
                gabaritos[numero] = respostas
    except Exception as e:
        logging.error(f"Erro ao ler gabarito: {e}")
        return {}
    
    return gabaritos

def corrigir_prova(nome, modelo, respostas, gabaritos):
    """Corrige uma prova individual baseada no gabarito."""
    try:
        gabarito = gabaritos.get(str(modelo))
        if not gabarito:
            logging.warning(f"Modelo {modelo} não encontrado no gabarito.")
            return 0

        respostas_aluno = respostas.replace("Respostas:", "").strip().split(",")
        respostas_aluno = [r.strip().upper() for r in respostas_aluno][:len(gabarito)]

        acertos = sum(1 for a, b in zip(gabarito, respostas_aluno) if a == b)
        return acertos

    except Exception as e:
        logging.error(f"Erro ao corrigir prova de {nome} (modelo {modelo}): {e}")
        return 0

@prova_bp.route('/verificar-email', methods=['POST'])
@cross_origin()
def verificar_email():
    """Verifica se o e-mail está autorizado."""
    try:
        data = request.get_json()
        email = data.get('email', '').strip().lower()
        
        if is_email_authorized(email):
            log_action(email, "LOGIN_AUTORIZADO", "Acesso concedido ao sistema")
            return jsonify({'autorizado': True, 'mensagem': 'Acesso autorizado'})
        else:
            log_action(email, "LOGIN_NEGADO", "Tentativa de acesso não autorizado")
            return jsonify({'autorizado': False, 'mensagem': 'Acesso não autorizado'})
    
    except Exception as e:
        log_action("DESCONHECIDO", "ERRO_LOGIN", error=str(e))
        return jsonify({'erro': 'Erro interno do servidor'}), 500


def resize_image(image, max_width=1000):
    """Redimensiona a imagem para largura máxima (mantendo proporção)."""
    w, h = image.size
    if w > max_width:
        ratio = max_width / w
        new_size = (max_width, int(h * ratio))
        return image.resize(new_size, Image.LANCZOS)
    return image


@prova_bp.route('/processar', methods=['POST'])
@cross_origin()
def processar_provas():
    """Processa os arquivos Excel e PDF enviados."""
    email = request.form.get('email', '').strip()

    try:
        if not GEMINI_API_KEY or not model:
            log_action(email, "ERRO_PROCESSAMENTO", error="API Gemini não configurada ou não disponível")
            return jsonify({'erro': 'API Gemini não configurada ou não disponível no ambiente atual'}), 500

        if not convert_from_path or not Image:
            log_action(email, "ERRO_DEPENDENCIAS", error="Dependências críticas não disponíveis")
            return jsonify({'erro': 'Dependências críticas não disponíveis no ambiente atual'}), 500

        if 'excel' not in request.files or 'pdf' not in request.files:
            log_action(email, "ERRO_UPLOAD", error="Arquivos Excel e PDF são obrigatórios")
            return jsonify({'erro': 'Arquivos Excel e PDF são obrigatórios'}), 400

        excel_file = request.files['excel']
        pdf_file = request.files['pdf']
        tipo_prova = request.form.get('tipo_prova', 'novo').lower()

        if not is_email_authorized(email):
            log_action(email, "ACESSO_NEGADO", error="E-mail não autorizado para processamento")
            return jsonify({'erro': 'E-mail não autorizado'}), 403

        log_action(email, "INICIO_PROCESSAMENTO", f"Arquivos: {excel_file.filename}, {pdf_file.filename}")

        temp_dir = tempfile.mkdtemp()
        try:
            excel_path = os.path.join(temp_dir, secure_filename(excel_file.filename))
            pdf_path = os.path.join(temp_dir, secure_filename(pdf_file.filename))
            excel_file.save(excel_path)
            pdf_file.save(pdf_path)

            gabaritos = carregar_gabaritos_excel(excel_path)
            if not gabaritos:
                log_action(email, "ERRO_GABARITO", error="Não foi possível carregar os gabaritos do Excel")
                return jsonify({'erro': 'Não foi possível carregar os gabaritos do Excel'}), 400

            log_action(email, "GABARITOS_CARREGADOS", f"Modelos encontrados: {list(gabaritos.keys())}")

            try:
                info = pdfinfo_from_path(pdf_path, userpw=None)
                total_paginas = int(info.get("Pages", 0))
                if total_paginas <= 0:
                    raise RuntimeError("PDF sem páginas detectadas")
                log_action(email, "PDF_INFO", f"Total de páginas: {total_paginas}")
            except Exception as e:
                log_action(email, "ERRO_PDF_INFO", error=f"Erro ao ler info do PDF: {str(e)}")
                return jsonify({'erro': 'Erro ao ler informações do PDF'}), 400

            wb = Workbook()
            ws = wb.active
            ws.title = "Resultados"
            ws.append(["Nome", "Modelo", "Nota"])
            provas_processadas = 0

            COORDS = COORDS_ANTIGO if tipo_prova == "antigo" else COORDS_NOVO
            DPI = 150  # menor para reduzir RAM

            import gc
            for page_index in range(1, total_paginas + 1):
                pagina = None
                try:
                    log_action(email, "PROCESSANDO_PAGINA", f"Página {page_index} de {total_paginas}")

                    pages = convert_from_path(
                        pdf_path,
                        dpi=DPI,
                        first_page=page_index,
                        last_page=page_index,
                        thread_count=1
                    )
                    pagina = pages[0]

                    # Cortes + resize
                    img_nome = resize_image(crop_pil(pagina, COORDS["BOX_NOME"]))
                    img_nome = preprocessar_para_ia(img_nome)
                    img_modelo = resize_image(crop_pil(pagina, COORDS["BOX_MODELO"]))
                    img_resposta = resize_image(crop_pil(pagina, COORDS["BOX_RESPOSTA"]))

                    # Nome
                    prompt_nome = ("Qual o nome completo do aluno nesta imagem? Mostre apenas o que está escrito."
                                   "não considere hifens nem pontuações, apenas letras normais"
                                   "Apresente o nome sempre com as iniciais maiúsculas e as demais minúsculas.")
                    resposta_nome = model.generate_content([prompt_nome, img_nome])
                    nome_texto = resposta_nome.text.strip()

                    # Modelo
                    prompt_modelo = "Qual é o modelo do gabarito nesta imagem (Modelo 1, 2...)? Apenas o número."
                    resposta_modelo = model.generate_content([prompt_modelo, img_modelo])
                    modelo_texto = resposta_modelo.text.strip()

                    # Respostas
                    prompt_resposta = (
                        "Liste as alternativas marcadas no cartão-resposta desta imagem.\n"
                        "Considere apenas A, B, C, D ou E.\n"
                        "Formato: 'Respostas: A, B, C...'"
                    )
                    resposta_resposta = model.generate_content([prompt_resposta, img_resposta])
                    respostas_texto = resposta_resposta.text.strip()

                    acertos = corrigir_prova(nome_texto, modelo_texto, respostas_texto, gabaritos)
                    ws.append([nome_texto, modelo_texto, acertos])

                    provas_processadas += 1
                    log_action(email, "PROVA_CORRIGIDA",
                               f"Página {page_index} - Nome: {nome_texto}, Modelo: {modelo_texto}, Nota: {acertos}")

                except Exception as e:
                    log_action(email, "ERRO_PAGINA", error=f"Erro ao processar página {page_index}: {str(e)}")
                finally:
                    # libera memória
                    try:
                        if pagina: pagina.close()
                        img_nome.close()
                        img_modelo.close()
                        img_resposta.close()
                        del pages
                    except:
                        pass
                    gc.collect()

            resultado_path = os.path.join(temp_dir, 'resultado_provas.xlsx')
            wb.save(resultado_path)

            log_action(email, "PROCESSAMENTO_CONCLUIDO", f"Total de provas processadas: {provas_processadas}")
            return send_file(resultado_path, as_attachment=True,
                             download_name='resultado_provas.xlsx',
                             mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        finally:
            def cleanup():
                try:
                    shutil.rmtree(temp_dir)
                    log_action(email, "LIMPEZA_CONCLUIDA")
                except Exception as e:
                    log_action(email, "ERRO_LIMPEZA", error=str(e))

            import threading
            threading.Timer(CLEANUP_DELAY_SECONDS, cleanup).start()

    except Exception as e:
        log_action(email, "ERRO_GERAL", error=f"Erro geral: {str(e)}")
        return jsonify({'erro': 'Erro interno do servidor'}), 500


@prova_bp.route('/status', methods=['GET'])
@cross_origin()
def status():
    """Retorna o status da API."""
    return jsonify({
        'status': 'online',
        'gemini_configurado': bool(GEMINI_API_KEY and model),
        'emails_autorizados_count': len(EMAILS_AUTORIZADOS),
        'timestamp': datetime.now().isoformat()
    })

@prova_bp.route('/logs', methods=['GET'])
@cross_origin()
def get_logs():
    """Retorna os últimos logs do sistema (apenas para administradores)."""
    try:
        # Em produção, você deveria implementar autenticação de administrador aqui
        if not os.path.exists(LOG_FILE):
            return jsonify({'logs': []})
        
        with open(LOG_FILE, 'r', encoding='utf-8') as f:
            lines = f.readlines()
        
        # Retornar apenas as últimas 100 linhas
        recent_logs = lines[-100:] if len(lines) > 100 else lines
        
        return jsonify({
            'logs': [line.strip() for line in recent_logs],
            'total_lines': len(lines)
        })
    
    except Exception as e:
        return jsonify({'erro': f'Erro ao ler logs: {str(e)}'}), 500

